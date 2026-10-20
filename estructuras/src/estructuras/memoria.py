"""Memoria de cálculo estructural en Word (.docx) para edificios de albañilería confinada."""

import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from . import __version__
from .albanileria import ResultadoAlbanileria
from .analisis import verificar_densidad
from .cimentacion import INCREMENTO_SISMO, ResultadoCimentacion
from .geometria import luz_libre_losa
from .metrado import OpcionesMetrado
from .normas import e030, e070, vigencia
from .predimensionamiento import espesor_aligerado, seccion_viga

ROJO = RGBColor(0xB0, 0x00, 0x20)
MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre")


def _f(x: float, d: int = 2) -> str:
    return f"{x:.{d}f}"


def _barras(texto: str) -> str:
    return texto.replace("%%c", "Ø")


class _Doc:
    def __init__(self):
        self.d = Document()
        estilo = self.d.styles["Normal"]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(10)

    def titulo(self, texto, nivel=1):
        self.d.add_heading(texto, level=nivel)

    def p(self, texto, negrita=False, color=None, centrado=False):
        par = self.d.add_paragraph()
        run = par.add_run(texto)
        run.bold = negrita
        if color:
            run.font.color.rgb = color
        if centrado:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return par

    def vinetas(self, items):
        for it in items:
            self.d.add_paragraph(it, style="List Bullet")

    def tabla(self, encabezado, filas, nota=None):
        t = self.d.add_table(rows=1, cols=len(encabezado))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for celda, texto in zip(t.rows[0].cells, encabezado):
            celda.text = str(texto)
            for r in celda.paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(8.5)
        for fila in filas:
            celdas = t.add_row().cells
            for celda, texto in zip(celdas, fila):
                celda.text = str(texto)
                for r in celda.paragraphs[0].runs:
                    r.font.size = Pt(8.5)
        if nota:
            par = self.d.add_paragraph()
            run = par.add_run(nota)
            run.italic = True
            run.font.size = Pt(8)
        else:
            self.d.add_paragraph()

    def guardar(self, ruta: Path):
        ruta.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(ruta)


def generar(res: ResultadoAlbanileria, cim: ResultadoCimentacion, ruta: Path) -> Path:
    f = res.ficha
    met = res.metrado
    unidad = e070.unidad(f.unidad_albanileria)
    sitio = e030.parametros_sitio(f.zona, f.perfil_suelo, f.vs30)
    op = OpcionesMetrado.desde_ficha(f)
    doc = _Doc()
    n = f.n_pisos

    # ---------------------------------------------------------------- portada
    doc.p("MEMORIA DE CÁLCULO ESTRUCTURAL", negrita=True, centrado=True)
    doc.p(f.nombre.upper(), negrita=True, centrado=True)
    ubic = ", ".join(v for v in (f.ubicacion.get("direccion"), f.ubicacion.get("distrito"),
                                 f.ubicacion.get("provincia"), f.ubicacion.get("departamento")) if v)
    doc.p(ubic or "Ubicación: por definir", centrado=True)
    hoy = datetime.date.today()
    doc.p(f"{MESES[hoy.month - 1].capitalize()} de {hoy.year}", centrado=True)
    doc.p("BORRADOR GENERADO CON EL PROGRAMA «estructuras» v" + __version__ +
          ". Requiere revisión, validación y firma del ingeniero responsable.", color=ROJO, centrado=True)

    # ---------------------------------------------------------------- 1. generalidades
    doc.titulo("1. Generalidades")
    doc.titulo("1.1 Alcance", 2)
    doc.p(f"Se presenta el análisis y diseño estructural de una edificación de {n} pisos de uso "
          f"{f.uso or 'vivienda'}, con sistema de albañilería confinada en ambas direcciones, según el "
          "Reglamento Nacional de Edificaciones (RNE).")
    doc.titulo("1.2 Normas empleadas", 2)
    filas = []
    for nrm in vigencia.NORMAS:
        if nrm.codigo in ("E.020", "E.030", "E.050", "E.060", "E.070"):
            filas.append([nrm.codigo, nrm.titulo, "; ".join(nrm.dispositivos)])
    doc.tabla(["Norma", "Título", "Dispositivo legal"], filas,
              f"Vigencia revisada al {vigencia.FECHA_REVISION}.")
    doc.titulo("1.3 Materiales", 2)
    doc.vinetas([
        f"Concreto armado: f'c = {f.fc:g} kg/cm² (columnas, vigas, soleras, losas y placas).",
        f"Acero de refuerzo: fy = {f.fy:g} kg/cm² (grado 60).",
        f"Albañilería: {unidad.nombre}, f'b = {unidad.fb:g} kg/cm², f'm = {unidad.fm:g} kg/cm², "
        f"v'm = {unidad.vm:g} kg/cm² (E.070 Tabla 9), Em = 500 f'm, Gm = 0,4 Em (art. 24.7).",
        "Mortero 1:4 (cemento:arena gruesa), juntas de 1,0 a 1,5 cm.",
    ])
    if n >= 3 and f.zona >= 3:
        doc.p("Para edificios de 3 a 5 pisos en zonas sísmicas altas, la E.070 (Tabla 7) exige obtener "
              "f'm y v'm mediante ensayos de pilas y muretes; los valores de la Tabla 9 deben "
              "confirmarse en obra.", color=ROJO)
    doc.titulo("1.4 Suelo y parámetros de sitio", 2)
    q = f"{f.capacidad_portante:g} kg/cm²" if f.capacidad_portante else "pendiente"
    doc.vinetas([
        f"Zona sísmica {f.zona} (Z = {_f(sitio.Z)}), perfil de suelo {sitio.perfil} "
        f"({e030.NOMBRE_PERFIL[sitio.perfil]}): S = {_f(sitio.S)}, TP = {_f(sitio.TP)} s, TL = {_f(sitio.TL)} s.",
        f"Capacidad portante admisible: {q}; profundidad de cimentación Df = {_f(cim.opciones.df)} m.",
    ])
    if f.capacidad_portante_supuesta:
        doc.p("La capacidad portante y el perfil de suelo son SUPUESTOS hasta contar con el Estudio de "
              "Mecánica de Suelos (E.050).", color=ROJO)

    # ---------------------------------------------------------------- 2. estructuración
    doc.titulo("2. Estructuración y predimensionamiento")
    doc.p("Los muros portantes son de albañilería confinada por columnas y vigas soleras de concreto "
          "armado (E.070 art. 20), continuos hasta la cimentación. En fachada se emplean placas de "
          "concreto armado. Los muros de un mismo eje unidos por columnas se analizan como un solo muro "
          "de varios paños (Nc columnas).")
    filas = [[w.id, w.direccion.upper(), "Albañilería" if w.material == "albanileria" else "Concreto",
              _f(w.L), _f(w.espesor), w.Nc, ", ".join(p.id for p in w.panos)] for w in res.muros]
    doc.tabla(["Muro", "Dir.", "Material", "L (m)", "t (m)", "Nc", "Paños"], filas)

    geo = f.geometria
    doc.titulo("2.1 Losas y vigas", 2)
    filas = []
    for i, losa in enumerate(geo.losas, start=1):
        luz = luz_libre_losa(losa)
        filas.append([i, losa.tipo, losa.direccion.upper(), _f(luz), _f(espesor_aligerado(luz)), _f(losa.h)])
    doc.tabla(["Paño", "Tipo", "Viguetas", "Luz (m)", "h ≥ L/25", "h adoptado"], filas)
    filas = []
    for v in geo.vigas:
        sv = seccion_viga(v.luz)
        filas.append([v.id, _f(v.luz), f"{_f(sv.h_min)} – {_f(sv.h_max)}", f"{_f(v.b)} × {_f(v.h)}"])
    doc.tabla(["Viga", "Luz (m)", "h = L/12 – L/10", "b × h adoptado (m)"], filas)

    doc.titulo("2.2 Densidad mínima de muros (E.070 art. 19.2b)", 2)
    dens = verificar_densidad(f)
    filas = [[d.eje, _f(d.area_muros, 3), _f(dens.area_planta), _f(d.densidad, 4), _f(d.requerida, 4),
              "Cumple" if d.cumple else "NO CUMPLE"] for d in dens.direcciones]
    doc.tabla(["Dir.", "ΣL·t (m²)", "Ap (m²)", "ΣL·t/Ap", "Z·U·S·N/56", "Resultado"], filas,
              f"Placas transformadas con Ec/Em = {_f(dens.Ec_sobre_Em)}; solo muros con L > 1,20 m (art. 17c).")

    # ---------------------------------------------------------------- 3. metrado
    doc.titulo("3. Metrado de cargas (E.020)")
    doc.vinetas([
        f"Peso de la albañilería: {op.peso_albanileria:g} tonf/m³; concreto armado: 2,4 tonf/m³; "
        f"tarrajeo de {op.tarrajeo * 100:g} cm por cara.",
        f"Piso terminado: {op.piso_terminado:g} tonf/m²; tabiquería repartida: {op.tabiqueria:g} tonf/m²; "
        f"acabados de azotea: {op.acabados_azotea:g} tonf/m².",
        f"Sobrecarga: {op.sobrecarga:g} tonf/m²; azotea: {op.sobrecarga_azotea:g} tonf/m².",
        "Las losas aligeradas descargan en los apoyos perpendiculares a las viguetas; los muros "
        "paralelos reciben una franja de 4 veces el espesor de la losa. Las vigas se consideran "
        "simplemente apoyadas. Las escaleras descargan en partes iguales en sus muros de apoyo.",
    ])
    filas = [[niv.nombre, _f(p), f"({_f(c[0])}, {_f(c[1])})"]
             for niv, p, c in zip(f.niveles, met.pesos_sismicos, met.centros_masa)]
    filas.append(["Total", _f(sum(met.pesos_sismicos)), ""])
    doc.tabla(["Nivel", "Peso sísmico P (tonf)", "Centro de masas (m)"], filas,
              f"P = carga muerta + {int(e030.fraccion_carga_viva(f.categoria) * 100)} % de la carga viva "
              "(azotea 25 %), E.030 art. 31.")

    # ---------------------------------------------------------------- 4. análisis sísmico
    doc.titulo("4. Análisis sísmico (E.030-2026)")
    c = res.sismo
    doc.vinetas([
        f"Z = {_f(c.Z)}, U = {_f(c.U)}, S = {_f(c.S)}, C = {_f(c.C)} (T = hn/CT = "
        f"{_f(f.altura_total)}/60 = {_f(f.altura_total / 60, 3)} s ≤ TP).",
        f"Sistema: albañilería confinada, R0 = 3; R = R0·Ia·Ip = {_f(res.R)}. Sismo severo según "
        "E.070 art. 22a; sismo moderado = 1/2 del severo.",
        f"Cortante basal: V = Z·U·C·S/R · P = {_f(c.coeficiente, 4)} × {_f(c.P)} = {_f(c.V)} tonf.",
    ])
    filas = [[niv.nombre, _f(h), _f(p), _f(fi), _f(ve)]
             for niv, h, p, fi, ve in zip(f.niveles, res.modelo.alturas, met.pesos_sismicos,
                                          res.fuerzas, res.cortantes_entrepiso)]
    doc.tabla(["Nivel", "h (m)", "P (tonf)", "Fi (tonf)", "VE entrepiso (tonf)"], filas,
              "Fuerzas del sismo severo distribuidas según E.030 art. 35.")
    doc.titulo("4.1 Modelo estructural", 2)
    doc.p("Modelo pseudo-tridimensional con diafragmas rígidos de 3 grados de libertad por piso "
          "(E.030 art. 30.3). Cada muro es un voladizo empotrado en la base (E.070 art. 24.5) con "
          "deformaciones de flexión y corte; su sección incluye las columnas de confinamiento "
          "transformadas con Ec/Em y el 25 % de los muros transversales o 6t (art. 24.6). Las fuerzas "
          "se aplican en el centro de masas con excentricidad accidental de 0,05 veces la dimensión "
          "perpendicular (art. 37) y se combinan 100 % + 30 % (art. 28.1 y 33.3).")
    filas = []
    for k, (T, part) in enumerate(zip(res.modelo.periodos[:3], res.modelo.participacion[:3]), start=1):
        filas.append([k, _f(T, 3), f"{part[0] * 100:.0f} %", f"{part[1] * 100:.0f} %"])
    doc.tabla(["Modo", "T (s)", "Masa X", "Masa Y"], filas, "Periodos del modelo (referenciales).")
    doc.titulo("4.2 Derivas y regularidad torsional", 2)
    filas = []
    for i, niv in enumerate(f.niveles):
        filas.append([niv.nombre, _f(res.derivas["X"][i], 4), _f(res.derivas["Y"][i], 4),
                      _f(res.relacion_torsion["X"][i]), _f(res.relacion_torsion["Y"][i])])
    lim = e030.sistema("albanileria").deriva_max
    doc.tabla(["Entrepiso", "Deriva X", "Deriva Y", "Δmax/Δprom X", "Δmax/Δprom Y"], filas,
              f"Derivas inelásticas = {_f(e030.factor_desplazamiento(res.R, f.regular))} × elásticas "
              f"(art. 50); límite {lim} (Tabla 14). El criterio de irregularidad torsional se aplica solo "
              f"si la deriva supera {lim / 2} (Tabla 12).")

    # ---------------------------------------------------------------- 5. diseño de muros
    doc.titulo("5. Diseño de muros de albañilería (E.070)")
    doc.p("Fuerzas del sismo moderado (Ve, Me) por muro y verificaciones: esfuerzo axial (19.1b), "
          "control de fisuración Ve ≤ 0,55 Vm (26.2) y resistencia al agrietamiento diagonal "
          "Vm = 0,5 v'm·α·t·L + 0,23 Pg (26.3).")
    filas = []
    for v in res.verificaciones:
        filas.append([v.muro, v.piso, _f(v.Pg, 1), _f(v.sigma_m, 1), _f(v.sigma_adm, 1), _f(v.Ve), _f(v.Me),
                      _f(v.alfa), _f(v.Vm), "Sí" if v.ok_fisuracion else "NO",
                      _f(v.Vu), "Sí" if v.agrietado else "No"])
    doc.tabla(["Muro", "Piso", "Pg", "σm", "σadm", "Ve", "Me", "α", "Vm", "Ve≤0,55Vm", "Vu", "Agriet."], filas,
              "Fuerzas en tonf y tonf·m; esfuerzos en tonf/m². Vu = Ve·(Vm1/Ve1), con 2 ≤ Vm1/Ve1 ≤ 3 (art. 27c). "
              "Se diseñan como agrietados el primer piso y los entrepisos con Vu ≥ Vm (27.2).")
    doc.titulo("5.1 Resistencia al corte del edificio (art. 26.4)", 2)
    filas = [[g.direccion, g.piso, _f(g.suma_Vm), _f(g.VE), "Cumple" if g.cumple else "NO CUMPLE",
              "Sí" if g.elastico else "No"] for g in res.globales]
    doc.tabla(["Dir.", "Piso", "ΣVm (tonf)", "VE (tonf)", "ΣVm ≥ VE", "ΣVm ≥ 3VE"], filas,
              "ΣVm incluye los muros confinados con L ≥ 1,20 m y el aporte Vc de las placas.")
    horizontales = [f"{v.muro}, piso {v.piso}: {v.refuerzo_horizontal}" for v in res.verificaciones
                    if v.refuerzo_horizontal]
    doc.titulo("5.2 Refuerzo horizontal (art. 27.1)", 2)
    if horizontales:
        doc.p("Llevan refuerzo horizontal continuo, cuantía ρ ≥ 0,001 (1 Ø 1/4\" cada 2 hiladas en "
              "muros de soga; cada hilada en muros de cabeza), anclado en las columnas:")
        doc.vinetas(horizontales)
    else:
        doc.p("No se requiere refuerzo horizontal.")

    # ---------------------------------------------------------------- 6. confinamientos
    doc.titulo("6. Diseño de elementos de confinamiento")
    doc.p("Columnas de muros agrietados con la Tabla 11 (art. 27.3): Vc, T y C; sección por corte-"
          "fricción Acf = Vc/(0,2 f'c φ) y por compresión An = As + (C/φ − As·fy)/(0,85 δ f'c); "
          f"refuerzo As = Vc/(fy μ φ) + T/(fy φ) ≥ 0,1 f'c Ac/fy, μ = 0,8; estribos s1 a s4. Pisos no "
          "agrietados con el art. 27.4. En las intersecciones se coloca el mayor refuerzo (art. 26.1a).")
    filas = []
    for tipo, d in res.tipos_columna.items():
        nodos = "; ".join(f"({_f(x, 1)}, {_f(y, 1)})" for x, y in d["nodos"])
        filas.append([tipo, f"{d['t'] / 100:.2f} × {d['d'] / 100:.2f}", _barras(d["barras"]),
                      _barras(d["estribos"]), nodos])
    doc.tabla(["Tipo", "t × d (m)", "Refuerzo", "Estribos", "Ubicación (x, y)"], filas)
    filas = []
    for c_ in res.columnas:
        if c_.piso == 1:
            filas.append([c_.muro, f"({_f(c_.posicion[0], 1)}, {_f(c_.posicion[1], 1)})",
                          "Extrema" if c_.extrema else "Interior", _f(c_.Pc), _f(c_.Vc), _f(c_.T), _f(c_.C),
                          _f(c_.As_req), f"{c_.t:.0f} × {c_.d:.0f}"])
    doc.tabla(["Muro", "Nudo", "Columna", "Pc", "Vc", "T", "C", "As req (cm²)", "t × d (cm)"], filas,
              "Fuerzas del primer piso (tonf).")
    filas = {}
    for s in res.soleras:
        actual = filas.get(s.muro)
        if actual is None or s.As_req > actual[1]:
            filas[s.muro] = (s.Ts, s.As_req, s.barras)
    t_losa = max(l.h for l in geo.losas)
    doc.tabla(["Muro", "Ts máx (tonf)", "As req (cm²)", "Refuerzo"],
              [[k, _f(v[0]), _f(v[1]), _barras(v[2])] for k, v in sorted(filas.items())],
              f"Vigas soleras de ancho t y peralte {t_losa:.2f} m; Ts = V·Lm/(2L) (art. 27.3b y 27.4d); "
              "estribos Ø 1/4\": 1@.05, 4@.10, r@.25.")
    if res.placas:
        doc.titulo("6.1 Placas de concreto armado (E.070 art. 25.1; E.060)", 2)
        filas = [[p.muro, p.piso, _f(p.Vu), _f(p.Mu), _f(p.phi_Vc), _f(p.rho_h, 4), _f(p.As_extremo)]
                 for p in res.placas]
        doc.tabla(["Placa", "Piso", "Vu (tonf)", "Mu (tonf·m)", "φVc (tonf)", "ρh", "As extremo (cm²)"], filas,
                  "Gravedad + sismo moderado. Verificación preliminar; el diseño por flexocompresión de "
                  "las placas se completa con la E.060.")

    # ---------------------------------------------------------------- 7. cimentación
    doc.titulo("7. Cimentación")
    doc.p(f"Capacidad admisible {_f(cim.q_admisible, 1)} tonf/m²; capacidad neta "
          f"q = qadm − γ·Df = {_f(cim.q_neta, 1)} tonf/m². Cimientos corridos de concreto ciclópeo "
          f"(1:10 + 30 % PG), peralte {_f(cim.opciones.h_cimiento)} m; sobrecimiento de concreto 1:8 + 25 % PM, "
          f"de ancho t y altura {_f(cim.opciones.h_sobrecimiento)} m.")
    filas = [[c_.elemento, c_.tipo, _f(c_.largo), _f(c_.ancho), _f(c_.P, 1), _f(c_.q_servicio, 1)]
             for c_ in cim.cimientos]
    doc.tabla(["Elemento", "Tipo", "Largo (m)", "Ancho (m)", "P = D + L (tonf)", "q (tonf/m²)"], filas)
    filas = [[v.direccion, _f(v.P, 1), _f(v.M, 1), _f(v.q_max, 1), _f(v.q_min, 1),
              _f(INCREMENTO_SISMO * cim.q_neta, 1)] for v in cim.sismo]
    doc.tabla(["Sismo", "P (tonf)", "M volteo (tonf·m)", "q máx", "q mín", "1,3 q neta"], filas,
              "Los cimientos forman una retícula continua que resiste el volteo en conjunto "
              "(sismo moderado, 100 % + 30 %).")

    # ---------------------------------------------------------------- 8. conclusiones
    doc.titulo("8. Conclusiones")
    if res.cumple and cim.cumple:
        doc.p("La estructura cumple las verificaciones de la E.030-2026 y de la E.070 consideradas en "
              "esta memoria: densidad de muros, esfuerzo axial, control de fisuración, resistencia al "
              "corte del edificio, derivas y regularidad, y presiones en la cimentación.")
    else:
        doc.p("La estructura NO cumple todas las verificaciones; ver las observaciones.", color=ROJO)
    obs = res.observaciones + cim.observaciones
    if obs:
        doc.p("Observaciones:", negrita=True)
        doc.vinetas(obs)
    doc.titulo("9. Recomendaciones y limitaciones")
    doc.vinetas([
        "Confirmar el perfil de suelo y la capacidad portante con el Estudio de Mecánica de Suelos.",
        "Verificar f'm y v'm con ensayos de pilas y muretes (E.070 art. 13).",
        "El modelo considera los muros en voladizo sin acoplamiento por vigas o losas, y no incluye el "
        "efecto de los alféizares; los alféizares y tabiques deben aislarse de los muros portantes.",
        "Pendiente: diseño de vigas, losas, escalera y placas por flexocompresión (E.060), y refuerzo de "
        "las zapatas de placas y columnas.",
        "La supervisión de obra debe verificar la calidad de la albañilería y el confinamiento.",
    ])
    ruta = Path(ruta)
    doc.guardar(ruta)
    return ruta
